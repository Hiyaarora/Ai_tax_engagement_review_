"""Generate the synthetic demo documents into data/synthetic/.

Usage (from backend/):  uv run python -m scripts.synthetic.generate [output_dir]

PDFs are kept to 2 pages because the Document Intelligence F0 tier only analyzes the first 2 pages.
"""

from __future__ import annotations

import io
import json
import sys
from dataclasses import asdict
from pathlib import Path
from typing import Any

from docx import Document
from reportlab.lib import colors
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from scripts.synthetic.dataset import (
    COMPANY,
    ENGAGEMENT_ID,
    TAX_YEAR,
    employee_locations,
    questionnaire,
    sales_transactions,
    write_sales_csv,
)

THRESHOLDS_PATH = Path(__file__).resolve().parents[2] / "app/tools/reference_data/thresholds.json"
DEFAULT_OUTPUT_DIR = Path(__file__).resolve().parents[3] / "data/synthetic"

ENGAGEMENT_DIR = f"engagements/{ENGAGEMENT_ID}"
SHARED_DIR = "shared"

# Relative to the output dir. Documents are indexed as evidence; .csv/.json feed the tools.
OUTPUT_FILES = (
    f"{ENGAGEMENT_DIR}/engagement.json",
    f"{ENGAGEMENT_DIR}/questionnaire.pdf",
    f"{ENGAGEMENT_DIR}/questionnaire.json",
    f"{ENGAGEMENT_DIR}/locations.docx",
    f"{ENGAGEMENT_DIR}/locations.json",
    f"{ENGAGEMENT_DIR}/sales.csv",
    f"{SHARED_DIR}/salt_reference_guide.pdf",
)

_styles = getSampleStyleSheet()
_body = _styles["BodyText"]
_small = ParagraphStyle("small", parent=_body, fontSize=8, leading=10, textColor=colors.grey)
_table_style = TableStyle(
    [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#dddddd")),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
    ]
)


def _pdf(path: Path, title: str, flowables: list[Any]) -> None:
    doc = SimpleDocTemplate(
        str(path), pagesize=LETTER, title=title, leftMargin=inch, rightMargin=inch
    )
    doc.build(
        [
            Paragraph(COMPANY.disclaimer, _small),
            Spacer(1, 6),
            Paragraph(title, _styles["Title"]),
            *flowables,
        ]
    )


def _wrap(text: str) -> Paragraph:
    return Paragraph(text, _body)


def write_questionnaire_pdf(path: Path) -> None:
    items = questionnaire()
    flow: list[Any] = [
        _wrap(
            f"<b>Client:</b> {COMPANY.name} &nbsp;&nbsp; <b>Tax year:</b> {TAX_YEAR} &nbsp;&nbsp; "
            f"<b>Engagement:</b> {ENGAGEMENT_ID}"
        ),
        _wrap(
            "Completed by client management. Answers are self-reported and have not been verified "
            "by the engagement team."
        ),
        Spacer(1, 10),
    ]
    for section in dict.fromkeys(item.section for item in items):
        flow.append(Paragraph(section, _styles["Heading2"]))
        rows = [["#", "Question", "Answer", "Client note"]]
        rows += [
            [item.id, _wrap(item.question), _wrap(item.answer), _wrap(item.note)]
            for item in items
            if item.section == section
        ]
        table = Table(rows, colWidths=[1.55 * inch, 2.35 * inch, 0.9 * inch, 1.7 * inch])
        table.setStyle(_table_style)
        flow += [table, Spacer(1, 8)]
    _pdf(path, f"State and Local Tax (SALT) Nexus Questionnaire - {TAX_YEAR}", flow)


def write_sales_csv_file(path: Path) -> None:
    buffer = io.StringIO()
    write_sales_csv(sales_transactions(), buffer)
    header = (
        f"# SYNTHETIC DATA - NOT A REAL CLIENT. {COMPANY.name}, tax year {TAX_YEAR}, "
        "ship-to state sales. Lines starting with # are comments.\n"
    )
    path.write_text(header + buffer.getvalue(), encoding="utf-8")


def write_engagement_json(path: Path) -> None:
    path.write_text(
        json.dumps(
            {
                "engagement_id": ENGAGEMENT_ID,
                "company_name": COMPANY.name,
                "home_state": COMPANY.home_state,
                "tax_year": TAX_YEAR,
                "disclaimer": COMPANY.disclaimer,
            },
            indent=2,
        ),
        encoding="utf-8",
    )


def write_questionnaire_json(path: Path) -> None:
    path.write_text(json.dumps([asdict(q) for q in questionnaire()], indent=2), encoding="utf-8")


def write_locations_json(path: Path) -> None:
    path.write_text(
        json.dumps([asdict(loc) for loc in employee_locations()], indent=2), encoding="utf-8"
    )


def write_locations_docx(path: Path) -> None:
    doc = Document()
    doc.add_paragraph(COMPANY.disclaimer)
    doc.add_heading(f"{COMPANY.name} - Employee and Office Locations ({TAX_YEAR})", level=1)
    doc.add_paragraph(
        "Prepared by HR from the payroll system as of December 31. Headcount counts W-2 employees "
        "only; contractors and third-party sites are listed for completeness with headcount 0."
    )
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    headers = ["City", "State", "Site type", "Headcount", "Note"]
    for cell, text in zip(table.rows[0].cells, headers, strict=True):
        cell.text = text
    for loc in employee_locations():
        cells = table.add_row().cells
        cells[0].text, cells[1].text, cells[2].text = loc.city, loc.state, loc.site_type
        cells[3].text, cells[4].text = str(loc.headcount), loc.note
    total = sum(loc.headcount for loc in employee_locations())
    doc.add_paragraph(f"Total W-2 headcount: {total}.")
    doc.save(str(path))


def write_reference_pdf(path: Path) -> None:
    data = json.loads(THRESHOLDS_PATH.read_text(encoding="utf-8"))
    flow: list[Any] = [
        _wrap(
            "<b>ILLUSTRATIVE ONLY.</b> The thresholds and rules below are invented for a software "
            "demo and are not a statement of any state's law. Consult current authoritative "
            "guidance for actual thresholds."
        ),
        Spacer(1, 8),
        Paragraph("1. Physical presence nexus", _styles["Heading2"]),
        _wrap(
            "A business generally has physical presence nexus in a state if it maintains an "
            "office, stores inventory in the state (including inventory held at a third-party "
            "fulfillment or 3PL warehouse), or has employees or agents working in the state. "
            "Remote employees "
            "working from home offices are commonly treated as creating physical presence. Regular "
            "in-state visits by sales representatives, including independent contractors acting on "
            "the company's behalf, may also create nexus."
        ),
        Paragraph("2. Economic nexus", _styles["Heading2"]),
        _wrap(
            "After the 2018 <i>Wayfair</i> decision, states may require remote sellers to collect "
            "sales tax once sales into the state exceed a threshold measured over the "
            f"{data['measurement_period']}. Some states use a sales threshold only; others combine "
            "sales with a transaction count using either an AND or an OR rule."
        ),
        Spacer(1, 6),
    ]
    rows = [["State", "Sales threshold (USD)", "Transaction threshold", "Rule"]]
    for code, s in data["states"].items():
        txn = "n/a" if s["transaction_threshold"] is None else str(s["transaction_threshold"])
        rows.append([f"{s['name']} ({code})", f"{s['sales_threshold']:,}", txn, s["rule"]])
    table = Table(rows, colWidths=[1.6 * inch, 1.7 * inch, 1.6 * inch, 1.6 * inch])
    table.setStyle(_table_style)
    flow += [
        table,
        Spacer(1, 8),
        Paragraph("3. Registration and marketplace facilitators", _styles["Heading2"]),
        _wrap(
            "Once nexus exists, the seller must register and begin collecting. Sales made through "
            "a marketplace that collects on the seller's behalf are usually excluded from the "
            "seller's "
            "own collection obligation but are typically still counted toward economic nexus "
            "thresholds. Holding inventory in a state through a 3PL does not remove the seller's "
            "registration obligation."
        ),
    ]
    _pdf(path, "SALT Nexus Reference Guide (Synthetic)", flow)


def generate_all(output_dir: Path) -> list[Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    writers = {
        "engagement.json": write_engagement_json,
        "questionnaire.pdf": write_questionnaire_pdf,
        "questionnaire.json": write_questionnaire_json,
        "locations.docx": write_locations_docx,
        "locations.json": write_locations_json,
        "sales.csv": write_sales_csv_file,
        "salt_reference_guide.pdf": write_reference_pdf,
    }
    written = []
    for relative in OUTPUT_FILES:
        path = output_dir / relative
        path.parent.mkdir(parents=True, exist_ok=True)
        writers[path.name](path)
        written.append(path)
    return written


if __name__ == "__main__":
    target = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_OUTPUT_DIR
    for p in generate_all(target):
        print(f"wrote {p}")
